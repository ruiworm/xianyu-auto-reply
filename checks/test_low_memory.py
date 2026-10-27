"""轻量部署回归：小连接池、冷启动依赖，以及按需加载后的真实文件处理。"""
from __future__ import annotations

from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, Mock, patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend-web"))
sys.path.insert(0, str(ROOT))


def run_python(source: str, directory: Path = ROOT):
    result = subprocess.run(
        [sys.executable, "-c", source], cwd=directory,
        capture_output=True, text=True, timeout=60,
    )
    if result.returncode:
        raise AssertionError(result.stdout + result.stderr)


class ResourceConfigTests(unittest.TestCase):
    def test_small_pools_can_be_set_through_environment(self):
        from common.core.config import BaseConfig

        with patch.dict(os.environ, {"DB_POOL_SIZE": "5", "REDIS_POOL_SIZE": "20"}):
            settings = BaseConfig(_env_file=None)
        self.assertEqual(settings.db_pool_size, 5)
        self.assertEqual(settings.redis_pool_size, 20)

    def test_invalid_pool_sizes_are_rejected(self):
        from common.core.config import BaseConfig
        from pydantic import ValidationError

        for field in ("db_pool_size", "REDIS_POOL_SIZE"):
            for value in (0, -1):
                with self.subTest(field=field, value=value):
                    with self.assertRaises(ValidationError):
                        BaseConfig(_env_file=None, **{field: value})

    def test_large_existing_pool_settings_remain_supported(self):
        from common.core.config import BaseConfig

        settings = BaseConfig(_env_file=None, db_pool_size=800, REDIS_POOL_SIZE=800)
        self.assertEqual(settings.db_pool_size, 800)
        self.assertEqual(settings.redis_pool_size, 800)

    def test_compose_overlay_preserves_data_and_credentials(self):
        environment = {key: value for key, value in os.environ.items()
                       if not key.startswith("LOW_MEMORY_")}

        def config(*files):
            args = ["docker", "compose", "--env-file", os.devnull]
            for file in files:
                args.extend(["-f", file])
            return json.loads(subprocess.check_output(
                args + ["config", "--format", "json"], cwd=ROOT, text=True, env=environment,
            ))

        base = config("docker-compose.yml")
        light = config("docker-compose.yml", "docker-compose.low-memory.yml")
        self.assertEqual(base["volumes"], light["volumes"])
        self.assertEqual(base["services"].keys(), light["services"].keys())
        for name, service in base["services"].items():
            for key in ("volumes", "networks", "ports", "healthcheck"):
                self.assertEqual(service.get(key), light["services"][name].get(key))
            for key, value in service.get("environment", {}).items():
                if key.startswith("MYSQL_") or key == "REDIS_PASSWORD":
                    self.assertEqual(value, light["services"][name]["environment"][key])
        for name in ("backend-web", "websocket", "scheduler"):
            env = light["services"][name]["environment"]
            self.assertEqual(env["DB_POOL_SIZE"], "5")
            self.assertEqual(env["DB_MAX_OVERFLOW"], "5")
            self.assertEqual(env["REDIS_POOL_SIZE"], "20")
        self.assertIn("--max-connections=100", light["services"]["mysql"]["command"])


class LazyImportTests(unittest.TestCase):
    def test_all_three_services_start_without_unused_heavy_libraries(self):
        for service in ("backend-web", "websocket", "scheduler"):
            with self.subTest(service=service):
                run_python("""
import main
import sys
for name in ('cv2', 'numpy', 'openpyxl', 'DrissionPage', 'patchright'):
    assert name not in sys.modules, f'{name} loaded during service startup'
""", ROOT / service)

    def test_reading_captcha_concurrency_does_not_load_browser_engines(self):
        run_python("""
import sys
from common.services.captcha.concurrency import get_browser_stats
assert isinstance(get_browser_stats(), dict)
for name in ('DrissionPage', 'playwright', 'patchright'):
    assert name not in sys.modules, name
""")

    def test_public_captcha_exports_keep_alias_and_singleton_identity(self):
        run_python("""
from importlib import import_module
import common.services.captcha as captcha
from common.services.captcha import RetryStrategyStats, strategy_stats
assert isinstance(strategy_stats, RetryStrategyStats)
assert strategy_stats is import_module('common.services.captcha.strategy_stats').strategy_stats
namespace = {}
exec('from common.services.captcha import *', namespace)
assert all(name in namespace for name in captcha.__all__)
assert namespace['SliderStealth'] is namespace['PlaywrightSliderService']
assert namespace['concurrency_manager'] is import_module('common.services.captcha.concurrency').concurrency_manager
for name in captcha.__all__:
    assert getattr(captcha, name) is namespace[name]
try:
    captcha.no_such_service
except AttributeError:
    pass
else:
    raise AssertionError('unknown export must raise AttributeError')
""")


async def response_bytes(response):
    return b"".join([chunk async for chunk in response.body_iterator])


class SpreadsheetTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.user = SimpleNamespace(id=42, role="MEMBER")

    async def test_keywords_export_import_roundtrip(self):
        from fastapi import UploadFile
        from app.api.routes.keywords import export_keywords, import_keywords

        account = SimpleNamespace(account_id="test-account")
        account_service = SimpleNamespace(get_account_for_user=AsyncMock(return_value=account))
        row = {"keyword": "测试", "item_id": "12345678901234567890", "reply": "自动回复"}
        service = SimpleNamespace(
            list_keywords=AsyncMock(return_value=[row]), replace_text_keywords=AsyncMock(),
        )
        response = await export_keywords("test-account", self.user, account_service, service)
        file = UploadFile(filename="keywords.xlsx", file=BytesIO(await response_bytes(response)))
        result = await import_keywords("test-account", file, self.user, account_service, service)
        self.assertTrue(result.success)
        imported = service.replace_text_keywords.await_args.args[1][0]
        for key, value in row.items():
            self.assertEqual(imported[key], value)
        account_service.get_account_for_user.assert_awaited_with(42, "test-account")

    async def test_blacklist_export_import_roundtrip(self):
        from fastapi import UploadFile
        from app.api.routes.blacklist import export_personal_blacklist, import_personal_blacklist

        item = SimpleNamespace(
            account_id="test-account", buyer_id="12345678901234567890", buyer_nick="买家",
            item_id="98765432109876543210", reason="测试原因", is_enabled=False,
        )
        service = SimpleNamespace(
            list_personal=AsyncMock(return_value=([item], 1)),
            create_personal=AsyncMock(return_value=[item]),
        )
        response = await export_personal_blacklist(self.user, service)
        file = UploadFile(filename="blacklist.xlsx", file=BytesIO(await response_bytes(response)))
        result = await import_personal_blacklist(file, self.user, service)
        self.assertTrue(result["success"])
        imported = service.create_personal.await_args.kwargs
        self.assertEqual(imported["buyer_ids"], [item.buyer_id])
        self.assertEqual(imported["item_id"], item.item_id)
        self.assertEqual(imported["owner_id"], 42)
        self.assertFalse(imported["is_enabled"])

    async def test_personal_addresses_export_import_roundtrip(self):
        from fastapi import UploadFile
        from app.api.routes import personal_addresses as route

        address = "浙江省杭州市测试地址123号"
        service = SimpleNamespace(
            list_all_for_owner=AsyncMock(return_value=[SimpleNamespace(address=address)]),
            upsert_many=AsyncMock(return_value={"created": 1, "updated": 0}),
        )
        with patch.object(route, "UserPublishAddressService", return_value=service):
            response = await route.export_personal_addresses(self.user, None)
            file = UploadFile(filename="addresses.xlsx", file=BytesIO(await response_bytes(response)))
            result = await route.import_personal_addresses(file, self.user, None)
        self.assertTrue(result.success)
        service.upsert_many.assert_awaited_once_with(42, [address])

    async def test_account_workbook_export_import_and_text_format(self):
        from openpyxl import Workbook, load_workbook
        from app.services.account_export_service import AccountExportService, _write_sheet
        from app.services.account_import_service import AccountImportService, _read_sheet_rows

        workbook = Workbook()
        _write_sheet(workbook, "文本数据", ["编号", "启用"], [["12345678901234567890", True]])
        self.assertEqual(_read_sheet_rows(workbook, "文本数据"),
                         [{"编号": "12345678901234567890", "启用": "是"}])
        self.assertEqual(workbook["文本数据"]["A2"].number_format, "@")
        exporter = AccountExportService(AsyncMock())
        exporter._get_accounts = AsyncMock(return_value=[])
        content = await exporter.export_accounts(42)
        exported = load_workbook(content, read_only=True)
        self.assertIn("账号基本信息", exported.sheetnames)
        self.assertIn("关键词规则", exported.sheetnames)
        exported.close()
        session = AsyncMock()
        session.execute.return_value = Mock()
        session.execute.return_value.all.return_value = []
        session.execute.return_value.scalars.return_value.all.return_value = []
        importer = AccountImportService(session, owner_id=42)
        importer._get_accounts_to_start = AsyncMock(return_value=[])
        result = await importer.import_accounts(content.getvalue())
        self.assertTrue(result["success"])
        self.assertEqual(result["data"]["errors"], [])
        self.assertEqual(result["data"]["inserted"], 0)


class VideoCoverTests(unittest.TestCase):
    def test_video_cover_still_works_after_lazy_import(self):
        import cv2
        import numpy as np
        from PIL import Image
        from common.services.xianyu_publish_video import _extract_video_cover_sync

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.avi"
            writer = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"MJPG"), 5, (64, 48))
            self.assertTrue(writer.isOpened())
            try:
                frame = np.zeros((48, 64, 3), dtype=np.uint8)
                frame[:, :, 1] = 200
                writer.write(frame)
                writer.write(frame)
            finally:
                writer.release()
            cover, width, height = _extract_video_cover_sync(path.read_bytes(), ".avi")
        self.assertEqual((width, height), (64, 48))
        with Image.open(BytesIO(cover)) as image:
            self.assertEqual(image.format, "JPEG")
            self.assertEqual(image.size, (64, 48))


if __name__ == "__main__":
    unittest.main()
